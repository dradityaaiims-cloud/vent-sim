"""Convert the project's simple markdown to a Word file. Usage: python3 scripts/md2docx.py in.md out.docx
Supports: # title, ## and ### headings, paragraphs, **bold**, *italic*, `code`, [text](url), bullets, numbered lists,
tables, images (![caption](path)), horizontal rules (skipped), checkbox bullets. Image paths are relative to the md file."""
import re, sys, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
base = os.path.dirname(os.path.abspath(src))
doc = Document()
sec = doc.sections[0]
for side in ("left_margin", "right_margin"): setattr(sec, side, Inches(1))
sec.top_margin = sec.bottom_margin = Inches(0.9)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5)):
    h = doc.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor(0x1F, 0x2D, 0x3D)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]*\]\([^)]+\))")

def add_inline(par, text, size=None):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"): r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = re.match(r"\[([^\]]*)\]\(([^)]+)\)", part).groups()
            r = par.add_run(f"{label} ({url})" if label and label != url else url)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: r = par.add_run(part[1:-1]); r.italic = True
        else: r = par.add_run(part)
        if size: r.font.size = Pt(size)

def para(text, style=None, size=None, align=None, space=6):
    p = doc.add_paragraph(style=style); add_inline(p, text, size)
    p.paragraph_format.space_after = Pt(space)
    if align: p.alignment = align
    return p

lines = open(src, encoding="utf-8").read().split("\n")
i, first_h1 = 0, True
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip() or ln.strip() == "---": i += 1; continue
    m = re.match(r"^!\[(.*)\]\((.*)\)$", ln.strip())
    if m:
        cap, path = m.groups(); full = os.path.join(base, path)
        if os.path.exists(full):
            doc.add_picture(full, width=Inches(6.4)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c = para(cap, size=9, space=10); c.runs[0].italic = True if c.runs else None
        else: para(f"[missing figure: {path}]")
        i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells): rows.append(cells)
            i += 1
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, r in enumerate(rows):
            for ci in range(ncol):
                cell = t.cell(ri, ci); cell.text = ""; p = cell.paragraphs[0]
                add_inline(p, r[ci] if ci < len(r) else "", size=8.5)
                p.paragraph_format.space_after = Pt(0)
                if ri == 0:
                    for run in p.runs: run.bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(4); continue
    if ln.startswith("# "):
        doc.add_paragraph(ln[2:].strip(), style="Title") if first_h1 else doc.add_heading(ln[2:].strip(), 1); first_h1 = False; i += 1; continue
    if ln.startswith("## "): doc.add_heading(ln[3:].strip(), 1); i += 1; continue
    if ln.startswith("### "): doc.add_heading(ln[4:].strip(), 2); i += 1; continue
    mb = re.match(r"^(\s*)- (\[[ x]\] )?(.*)$", ln)
    if mb:
        txt = (("☐ " if mb.group(2) == "[ ] " else "☑ ") if mb.group(2) else "") + mb.group(3)
        p = para(txt, style="List Bullet", space=2)
        if len(mb.group(1)) >= 2: p.paragraph_format.left_indent = Inches(0.6)
        i += 1; continue
    mn = re.match(r"^(\d+)\. (.*)$", ln)
    if mn:
        p = para(f"{mn.group(1)}. {mn.group(2)}", space=2); p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.25)
        i += 1; continue
    para(ln); i += 1
doc.save(dst); print("wrote", dst)
